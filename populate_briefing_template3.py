import os
import sys
from datetime import datetime, timedelta

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

prefix = os.environ.get("MAIN_PATH", ".")

if "DATE_STR" in os.environ:
    date_str = os.environ["DATE_STR"]
else:
    date_str = (datetime.today() - timedelta(days=2)).strftime("%Y-%m-%d")

briefing_plots_dir = os.path.join(prefix, "plots", "briefing", date_str)
template_path = os.path.join(prefix, "WeatherbriefingKenya_template3.pptx")
output_path = os.path.join(prefix, f"s2s_briefing3_{date_str}.pptx")

# Map plot stem (filename without .png) -> full path. A shape is matched to a
# plot by its exact name, minus a trailing ".png" if present.
available = {}
for fname in os.listdir(briefing_plots_dir):
    if fname.lower().endswith(".png"):
        available[fname[:-4]] = os.path.join(briefing_plots_dir, fname)


def replace_picture(slide, shape, image_path):
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    shape._element.getparent().remove(shape._element)
    slide.shapes.add_picture(image_path, left, top, width, height)


prs = Presentation(template_path)

matched = set()
for slide in prs.slides:
    for shape in list(slide.shapes):
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        name = shape.name[:-4] if shape.name.lower().endswith(".png") else shape.name
        if name in available:
            replace_picture(slide, shape, available[name])
            matched.add(name)

prs.save(output_path)

unused = sorted(set(available) - matched)
print(f"Filled {len(matched)} picture shape(s) from {briefing_plots_dir} -> {output_path}")
if unused:
    print(
        f"WARNING: no matching shape for {len(unused)} plot(s): {', '.join(unused)}",
        file=sys.stderr,
    )
